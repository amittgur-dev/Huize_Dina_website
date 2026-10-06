"""Build the stimulus/comparison spreadsheet in the researcher's notation.

Notation: B = body (main part), E = edges (endpoint components);
          ext = extension (length doubled, thickness unchanged),
          enl = enlargement (every dimension doubled); no suffix = unchanged.
Condition codes in data/ map to: A -> "B E", L -> "Bext E", P -> "Benl Eenl",
E -> "B Eenl" (or "B Eext" where the spec says the edge is lengthened only).

Usage: python3 tools/build_notation_sheet.py   (from experiments/line-similarity)
Writes data/stimuli-notation.xlsx, data/stimuli-notation.csv, data/comparisons-notation.csv
"""
from pathlib import Path
import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
measurements = {r['asset_id']: r for r in csv.DictReader((root / 'data/measurements.csv').open())}
trials = json.loads((root / 'data/trials.json').read_text())
families = {f['id']: f for f in json.loads((root / 'data/families.json').read_text())}

BODY = {0: 'plain line', 1: 'shaft', 2: 'shaft', 3: 'black shaft', 4: 'shaft', 5: 'shaft', 6: 'snake body', 7: 'windowed carriage', 8: 'lamp pole'}
EDGES = {0: None, 1: 'two arrowheads', 2: 'two circles', 3: 'two red end segments', 4: 'two vertical end lines', 5: 'one abstract right endpoint',
         6: 'head (and tail taper)', 7: 'locomotive, connector, rear projection, wheels', 8: 'upper fixture and base'}
# Edge operation in the E condition. The specification states that for the red
# segments and the vertical end lines E lengthens the edge at unchanged
# thickness (an extension); for the arrow, circles and abstract endpoint the
# whole endpoint bounding box doubles (an enlargement).
E_EDGE_OP = {1: 'enl', 2: 'enl', 3: 'ext', 4: 'ext', 5: 'enl'}

def ops(family, variant):
    """(body operation, edge operation) for a variant; '' = unchanged, None = no edges."""
    has_edges = EDGES[family] is not None
    if variant == 'A': return '', '' if has_edges else None
    if variant == 'L': return 'ext', '' if has_edges else None
    if variant == 'P': return 'enl', 'enl' if has_edges else None
    if variant == 'E': return '', E_EDGE_OP[family]
    raise ValueError(variant)

def notation(family, variant):
    b, e = ops(family, variant)
    parts = ['B' + b]
    if e is not None: parts.append('E' + e)
    return ' '.join(parts)

NOTES = {
    '7-L': 'Window count 14 -> 28; window size, gaps, rear projection, connector, locomotive and wheel radius unchanged.',
    '7-P': 'Every dimension doubled; still 14 windows.',
    '8-A': 'Vertical object: body length is the pole height.', '8-L': 'Pole height doubled; fixture and base unchanged.', '8-P': 'Pole, fixture and base all doubled.',
    '6-L': 'Body length doubled; head unchanged.', '6-P': 'Body and head doubled.',
    '3-E': 'Each red segment 1.8 -> 3.6 mm long at the shared 0.9 mm thickness (edge lengthened, not enlarged).',
    '4-E': 'Vertical end line 3.6 -> 7.2 mm long at 0.9 mm thickness (edge lengthened, not enlarged).',
    '5-A': 'Only the right endpoint exists in this family.',
    '0-A': 'Control: no edge component.', '0-L': 'Control: no edge component.', '0-P': 'Control: no edge component.',
}

stim_cols = ['asset_id', 'family', 'family_name', 'group', 'condition_code', 'notation', 'body_operation', 'edge_operation',
             'body_part', 'edge_part', 'body_length_mm', 'body_thickness_mm', 'body_thickness_to_length', 'edge_width_mm', 'edge_height_mm',
             'total_width_mm', 'total_height_mm', 'svg', 'notes']
stim_rows = []
order = {'A': 0, 'E': 1, 'L': 2, 'P': 3}
for aid, m in sorted(measurements.items(), key=lambda kv: (int(kv[1]['family']), order[kv[1]['variant']])):
    fam = int(m['family']); v = m['variant']
    b, e = ops(fam, v)
    stim_rows.append({
        'asset_id': aid, 'family': fam, 'family_name': m['name'], 'group': families[fam]['group'], 'condition_code': v,
        'notation': notation(fam, v), 'body_operation': b or 'unchanged', 'edge_operation': ('n/a' if e is None else (e or 'unchanged')),
        'body_part': BODY[fam], 'edge_part': EDGES[fam] or 'none',
        'body_length_mm': float(m['main_length_mm']), 'body_thickness_mm': float(m['main_thickness_mm']), 'body_thickness_to_length': m['main_thickness_to_length'],
        'edge_width_mm': float(m['edge_width_mm']) if m['edge_width_mm'] else None, 'edge_height_mm': float(m['edge_height_mm']) if m['edge_height_mm'] else None,
        'total_width_mm': float(m['total_width_mm']), 'total_height_mm': float(m['total_height_mm']),
        'svg': Path(m['src']).name, 'notes': NOTES.get(aid, ''),
    })

comp_cols = ['trial_id', 'family', 'family_name', 'group', 'reference', 'comparison_1_code', 'comparison_1', 'comparison_2_code', 'comparison_2', 'comparison',
             'comparison_1_asset', 'comparison_2_asset', 'canonical_left', 'canonical_right']
comp_rows = []
for t in trials:
    fam = t['family']
    comp_rows.append({
        'trial_id': t['id'], 'family': fam, 'family_name': t['name'], 'group': t['group'], 'reference': notation(fam, 'A'),
        'comparison_1_code': t['left'], 'comparison_1': notation(fam, t['left']), 'comparison_2_code': t['right'], 'comparison_2': notation(fam, t['right']),
        'comparison': f"{notation(fam, t['left'])}  vs  {notation(fam, t['right'])}",
        'comparison_1_asset': f"{fam}-{t['left']}", 'comparison_2_asset': f"{fam}-{t['right']}",
        'canonical_left': t['left'], 'canonical_right': t['right'],
    })

# CSV copies
with (root / 'data/stimuli-notation.csv').open('w', newline='') as f:
    w = csv.DictWriter(f, stim_cols); w.writeheader(); w.writerows({k: ('' if v is None else v) for k, v in r.items()} for r in stim_rows)
with (root / 'data/comparisons-notation.csv').open('w', newline='') as f:
    w = csv.DictWriter(f, comp_cols); w.writeheader(); w.writerows(comp_rows)

# Workbook
FONT = 'Arial'
head_font = Font(name=FONT, bold=True, color='FFFFFF'); head_fill = PatternFill('solid', fgColor='000000')
body_font = Font(name=FONT); bold = Font(name=FONT, bold=True); mono = Font(name=FONT, bold=True)
thin = Side(style='thin', color='BBBBBB'); border = Border(bottom=thin)
wb = Workbook()

def write_table(ws, cols, rows, widths, highlight=()):
    ws.append(cols)
    for c in ws[1]:
        c.font = head_font; c.fill = head_fill; c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in rows:
        ws.append([r[c] for c in cols])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body_font; c.border = border; c.alignment = Alignment(vertical='top', wrap_text=True)
            if cols[c.column - 1] in highlight: c.font = mono
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

ws = wb.active; ws.title = 'Notation'
legend = [
    ('Symbol', 'Meaning'),
    ('B', 'Body: the main part of the object (line or shaft, snake body, train carriage, lamp pole).'),
    ('E', 'Edges: the endpoint components (arrowheads, circles, red end segments, vertical end lines, abstract endpoint, snake head, locomotive/rear parts, lamp fixture and base).'),
    ('ext', 'Extension: length doubled, thickness unchanged (e.g. shaft 36 -> 72 mm at 0.9 mm).'),
    ('enl', 'Enlargement: every dimension doubled (e.g. shaft 36 x 0.9 -> 72 x 1.8 mm; endpoint box 4.5 -> 9 mm).'),
    ('no suffix', 'Unchanged relative to the reference A.'),
    ('', ''),
    ('Condition code (data files)', 'Notation'),
    ('A  reference', 'B E   (control family 0 has no edges: B)'),
    ('E  endpoint modification only', 'B Eenl for arrow, circles and abstract endpoint; B Eext for red segments and vertical end lines, whose edges are lengthened at unchanged thickness (see STIMULUS_SPECIFICATION.md).'),
    ('L  main-part extension', 'Bext E   (control: Bext)'),
    ('P  proportional enlargement', 'Benl Eenl   (control: Benl)'),
    ('', ''),
    ('Comparison types', 'B Eenl vs Bext E (E/L) · B Eenl vs Benl Eenl (E/P) · Bext E vs Benl Eenl (L/P). Families 1-5 have all three; the control and the snake, train and lamp have L/P only.'),
    ('', ''),
    ('Source', 'Dimensions from data/measurements.csv and docs/STIMULUS_SPECIFICATION.md (handoff of 2026-10-05, stimulus set 2026-10-05-compact-control-v1). All sizes are physical millimetres after card calibration.'),
    ('Left/right', 'The Comparisons sheet lists the canonical order from data/trials.json. In the experiment, left/right placement is assigned at random per presentation and recorded (left_condition / right_condition).'),
]
for r in legend: ws.append(list(r))
for row in ws.iter_rows():
    for c in row:
        c.font = body_font; c.alignment = Alignment(vertical='top', wrap_text=True)
for r in (1, 8): 
    for c in ws[r]: c.font = head_font; c.fill = head_fill
for r in (2, 3, 4, 5, 6, 9, 10, 11, 12): ws.cell(row=r, column=1).font = bold
ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 110

write_table(wb.create_sheet('Stimuli'), stim_cols, stim_rows,
            [9, 7, 22, 14, 10, 11, 11, 11, 18, 30, 10, 10, 10, 10, 10, 10, 10, 10, 60], highlight=('notation',))
write_table(wb.create_sheet('Comparisons'), comp_cols, comp_rows,
            [8, 7, 22, 14, 11, 10, 11, 10, 11, 26, 11, 11, 10, 10], highlight=('reference', 'comparison_1', 'comparison_2', 'comparison'))
wb.save(root / 'data/stimuli-notation.xlsx')
print(f"Wrote {len(stim_rows)} stimuli and {len(comp_rows)} comparisons to data/stimuli-notation.xlsx (+ CSV copies)")
